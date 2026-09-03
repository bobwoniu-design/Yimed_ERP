import csv
import io
import json
import mimetypes
import os
import re
import shutil
import zipfile
from xml.etree import ElementTree

import frappe
from frappe import _
from frappe.utils import cint, get_datetime, now_datetime


MAX_INDEX_FILE_BYTES = 25 * 1024 * 1024
MAX_ARCHIVE_UNCOMPRESSED_BYTES = 100 * 1024 * 1024
MAX_ARCHIVE_ENTRIES = 10000
MAX_XML_MEMBER_BYTES = 10 * 1024 * 1024
MAX_TEXT_CHARS = 1_000_000
MAX_PDF_PAGES = 500
MAX_SPREADSHEET_CELLS = 200_000


class UnsupportedContent(Exception):
	pass


def _safe_error(code):
	return {
		"file-too-large": _("文件超过全文索引大小上限"),
		"archive-unsafe": _("压缩文档结构不安全或超过索引限制"),
		"unsupported": _("当前文件格式不支持全文提取"),
		"ocr-unavailable": _("本机未配置 OCR 引擎，图片已安全跳过"),
		"encrypted-pdf": _("加密 PDF 无法建立全文索引"),
		"parse-failed": _("文件正文提取失败"),
	}.get(code, _("文件正文提取失败"))


def _normalize(text):
	text = re.sub(r"\s+", " ", text or "").strip()
	return text[:MAX_TEXT_CHARS], text.casefold()[:MAX_TEXT_CHARS]


def _language(text):
	chinese = len(re.findall(r"[\u3400-\u9fff]", text))
	latin = len(re.findall(r"[A-Za-z]", text))
	if chinese and latin:
		return "mixed"
	if chinese:
		return "zh"
	if latin:
		return "en"
	return "und"


def _safe_zip(content):
	try:
		archive = zipfile.ZipFile(io.BytesIO(content))
	except (zipfile.BadZipFile, OSError):
		raise UnsupportedContent("archive-unsafe")
	infos = archive.infolist()
	if len(infos) > MAX_ARCHIVE_ENTRIES:
		raise UnsupportedContent("archive-unsafe")
	total = 0
	for info in infos:
		name = info.filename.replace("\\", "/")
		if name.startswith("/") or ".." in name.split("/"):
			raise UnsupportedContent("archive-unsafe")
		total += info.file_size
		if total > MAX_ARCHIVE_UNCOMPRESSED_BYTES:
			raise UnsupportedContent("archive-unsafe")
		if info.compress_size and info.file_size / info.compress_size > 200:
			raise UnsupportedContent("archive-unsafe")
	return archive


def _safe_xml_text(archive, member):
	info = archive.getinfo(member)
	if info.file_size > MAX_XML_MEMBER_BYTES:
		raise UnsupportedContent("archive-unsafe")
	data = archive.read(member)
	upper = data[:4096].upper()
	if b"<!DOCTYPE" in upper or b"<!ENTITY" in upper:
		raise UnsupportedContent("archive-unsafe")
	try:
		root = ElementTree.fromstring(data)
	except ElementTree.ParseError:
		raise UnsupportedContent("parse-failed")
	return " ".join(value.strip() for value in root.itertext() if value and value.strip())


def _extract_plain(content):
	for encoding in ("utf-8-sig", "utf-16", "gb18030"):
		try:
			return content.decode(encoding), 1
		except (UnicodeDecodeError, LookupError):
			continue
	raise UnsupportedContent("parse-failed")


def _extract_pdf(content):
	from pypdf import PdfReader

	try:
		reader = PdfReader(io.BytesIO(content))
		if reader.is_encrypted:
			raise UnsupportedContent("encrypted-pdf")
		if len(reader.pages) > MAX_PDF_PAGES:
			raise UnsupportedContent("archive-unsafe")
		parts = []
		for page in reader.pages:
			parts.append(page.extract_text() or "")
			if sum(len(part) for part in parts) >= MAX_TEXT_CHARS:
				break
		return "\n".join(parts), len(reader.pages)
	except UnsupportedContent:
		raise
	except Exception:
		raise UnsupportedContent("parse-failed")


def _extract_docx(content):
	archive = _safe_zip(content)
	members = [
		name for name in archive.namelist()
		if name == "word/document.xml" or re.fullmatch(r"word/(header|footer)\d+\.xml", name)
	]
	if "word/document.xml" not in members:
		raise UnsupportedContent("parse-failed")
	return "\n".join(_safe_xml_text(archive, name) for name in members), 1


def _extract_pptx(content):
	archive = _safe_zip(content)
	members = sorted(
		(name for name in archive.namelist() if re.fullmatch(r"ppt/slides/slide\d+\.xml", name)),
		key=lambda name: int(re.search(r"(\d+)", name).group(1)),
	)
	if not members or len(members) > MAX_PDF_PAGES:
		raise UnsupportedContent("parse-failed")
	return "\n".join(_safe_xml_text(archive, name) for name in members), len(members)


def _extract_xlsx(content):
	_safe_zip(content)
	from openpyxl import load_workbook

	try:
		workbook = load_workbook(io.BytesIO(content), read_only=True, data_only=True, keep_links=False)
		parts = []
		cells = 0
		for sheet in workbook.worksheets[:100]:
			parts.append(sheet.title)
			for row in sheet.iter_rows(values_only=True):
				values = [str(value) for value in row if value not in (None, "")]
				cells += len(row)
				if cells > MAX_SPREADSHEET_CELLS:
					raise UnsupportedContent("archive-unsafe")
				if values:
					parts.append(" | ".join(values))
				if sum(len(part) for part in parts) >= MAX_TEXT_CHARS:
					break
		return "\n".join(parts), len(workbook.worksheets)
	except UnsupportedContent:
		raise
	except Exception:
		raise UnsupportedContent("parse-failed")


def _extract_image_ocr(content):
	if not shutil.which("tesseract"):
		raise UnsupportedContent("ocr-unavailable")
	try:
		import pytesseract
		from PIL import Image
	except ImportError:
		raise UnsupportedContent("ocr-unavailable")
	try:
		image = Image.open(io.BytesIO(content))
		image.verify()
		image = Image.open(io.BytesIO(content))
		if image.width * image.height > 50_000_000:
			raise UnsupportedContent("archive-unsafe")
		return pytesseract.image_to_string(image, timeout=30), 1
	except UnsupportedContent:
		raise
	except Exception:
		raise UnsupportedContent("parse-failed")


def _extract(content, filename, mime_type):
	extension = os.path.splitext((filename or "").lower())[1]
	if extension in {".txt", ".csv", ".json", ".md", ".log"}:
		return _extract_plain(content)
	if extension == ".pdf" or mime_type == "application/pdf":
		return _extract_pdf(content)
	if extension == ".docx":
		return _extract_docx(content)
	if extension == ".xlsx":
		return _extract_xlsx(content)
	if extension == ".pptx":
		return _extract_pptx(content)
	if extension in {".png", ".jpg", ".jpeg", ".tif", ".tiff", ".webp"}:
		return _extract_image_ocr(content)
	raise UnsupportedContent("unsupported")


def _index_doc(asset_doc, version_doc):
	name = frappe.db.get_value("Channel Content Search Index", {"asset": asset_doc.name}, "name")
	values = {
		"asset": asset_doc.name, "company": asset_doc.company, "folder": asset_doc.folder,
		"version": version_doc.name, "status": "Pending", "file_type": asset_doc.file_type,
		"mime_type": version_doc.mime_type or asset_doc.mime_type, "file_size": version_doc.size_bytes or asset_doc.file_size,
		"extracted_text": None, "normalized_text": None, "language": None, "page_count": 0,
		"char_count": 0, "error_summary": None, "queued_at": now_datetime(), "indexed_at": None,
	}
	if name:
		frappe.db.set_value("Channel Content Search Index", name, values)
		return frappe.get_doc("Channel Content Search Index", name)
	return frappe.get_doc({"doctype": "Channel Content Search Index", **values}).insert(ignore_permissions=True)


def queue_asset_index(asset, force=False, now=False):
	from channel_erp.content_center import append_content_audit
	from channel_erp.content_insights import invalidate_content_insights

	asset_doc = frappe.get_doc("Channel Content Asset", asset)
	if asset_doc.status != "Active" or not asset_doc.current_version:
		return {"asset": asset, "queued": False, "reason": "inactive-or-no-version"}
	version_doc = frappe.get_doc("Channel Content Asset Version", asset_doc.current_version)
	existing = frappe.db.get_value(
		"Channel Content Search Index", {"asset": asset}, ["version", "status"], as_dict=True
	)
	if existing and existing.version == version_doc.name and existing.status in {"Pending", "Processing", "Ready"} and not force:
		return {"asset": asset, "version": version_doc.name, "queued": False, "status": existing.status}
	index_doc = _index_doc(asset_doc, version_doc)
	invalidate_content_insights(asset)
	append_content_audit("index-enqueue", asset_doc, details={"version": version_doc.name, "force": bool(force)})
	if now:
		process_content_index(asset, version_doc.name)
	else:
		frappe.enqueue(
			"channel_erp.content_index.process_content_index", queue="short", timeout=300,
			asset=asset, version=version_doc.name, enqueue_after_commit=True,
			job_id=f"content-index-{asset}-{version_doc.name}", deduplicate=True,
		)
	return {"asset": asset, "version": version_doc.name, "index": index_doc.name, "queued": True}


def process_content_index(asset, version):
	from channel_erp.content_center import _file_content_bytes, append_content_audit

	if not frappe.db.exists("Channel Content Asset", asset):
		return
	asset_doc = frappe.get_doc("Channel Content Asset", asset)
	if asset_doc.status != "Active" or asset_doc.current_version != version:
		return
	index_name = frappe.db.get_value("Channel Content Search Index", {"asset": asset}, "name")
	if not index_name:
		return
	frappe.db.set_value("Channel Content Search Index", index_name, "status", "Processing")
	try:
		version_doc = frappe.get_doc("Channel Content Asset Version", version)
		frappe.clear_document_cache("File", version_doc.file)
		file_doc = frappe.get_doc("File", version_doc.file)
		file_doc.reload()
		file_size = cint(file_doc.file_size or version_doc.size_bytes)
		if file_size > MAX_INDEX_FILE_BYTES:
			raise UnsupportedContent("file-too-large")
		content = _file_content_bytes(file_doc)
		if len(content) > MAX_INDEX_FILE_BYTES:
			raise UnsupportedContent("file-too-large")
		text, page_count = _extract(content, version_doc.original_filename or file_doc.file_name, version_doc.mime_type)
		extracted, normalized = _normalize(text)
		if asset_doc.current_version != version or asset_doc.status != "Active":
			return
		frappe.db.set_value(
			"Channel Content Search Index", index_name,
			{"status": "Ready", "extracted_text": extracted, "normalized_text": normalized,
			 "language": _language(extracted), "page_count": page_count, "char_count": len(extracted),
			 "error_summary": None, "indexed_at": now_datetime()},
		)
		append_content_audit("index-ready", asset_doc, details={"version": version, "char_count": len(extracted)})
		from channel_erp.content_insights import queue_content_insights
		queue_content_insights(asset, force=True, requested_by=frappe.session.user)
	except UnsupportedContent as exc:
		code = str(exc) or "unsupported"
		frappe.db.set_value(
			"Channel Content Search Index", index_name,
			{"status": "Unsupported", "error_summary": _safe_error(code), "indexed_at": now_datetime()},
		)
		append_content_audit("index-unsupported", asset_doc, details={"version": version, "reason": code})
	except Exception:
		frappe.db.set_value(
			"Channel Content Search Index", index_name,
			{"status": "Failed", "error_summary": _safe_error("parse-failed"), "indexed_at": now_datetime()},
		)
		append_content_audit("index-failed", asset_doc, details={"version": version, "reason": "parse-failed"})


def invalidate_asset_index(asset, delete=False):
	from channel_erp.content_insights import invalidate_content_insights
	invalidate_content_insights(asset, delete=delete)
	name = frappe.db.get_value("Channel Content Search Index", {"asset": asset}, "name")
	if not name:
		return
	if delete:
		doc = frappe.get_doc("Channel Content Search Index", name)
		doc.flags.allow_index_delete = True
		doc.delete(ignore_permissions=True)
	else:
		frappe.db.set_value(
			"Channel Content Search Index", name,
			{"status": "Pending", "extracted_text": None, "normalized_text": None,
			 "char_count": 0, "error_summary": _("等待文件恢复后重建"), "indexed_at": None},
		)


def run_content_index_maintenance(limit=100):
	queued = 0
	for row in frappe.get_all(
		"Channel Content Asset", filters={"status": "Active"}, fields=["name", "current_version"], limit=cint(limit) or 100
	):
		index = frappe.db.get_value(
			"Channel Content Search Index", {"asset": row.name}, ["version", "status"], as_dict=True
		)
		if row.current_version and (not index or index.version != row.current_version or index.status in {"Failed", "Pending"}):
			if queue_asset_index(row.name).get("queued"):
				queued += 1
	return {"queued": queued}
