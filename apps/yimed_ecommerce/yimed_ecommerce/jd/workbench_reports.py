"""Read-only A4 work sheets and Excel exports for the procurement workbench."""
from collections import defaultdict
from io import BytesIO
import math

import frappe
from frappe import _
from frappe.utils import flt, now_datetime, format_datetime

from yimed_ecommerce.jd import workflow
from yimed_ecommerce.jd.report_pdf import _validate_batch_access, _render_chrome_pdf, _format_cell

TITLES = {"stocking": "商品备货汇总表", "assembly": "组合装组装单", "sorting": "采购订单分拣单"}


def column(key, label, width=90, numeric=False):
	return {"fieldname": key, "label": label, "width": width, "fieldtype": "Float" if numeric else "Data"}


def _draft_rows(value):
	if value is None:
		return None
	rows = frappe.parse_json(value) if isinstance(value, str) else value
	if not isinstance(rows, list) or len(rows) > 10000 or any(not isinstance(row, dict) for row in rows):
		frappe.throw(_("导出明细格式无效。"))
	for row in rows:
		for key in ("stocked_qty", "sorted_qty"):
			if key in row and (not math.isfinite(flt(row[key])) or flt(row[key]) < 0):
				frappe.throw(_("导出数量必须为非负数。"))
	return rows


def build_report(import_batch, stage, purchase_order=None, draft_rows=None, dimension="order"):
	if stage not in TITLES:
		frappe.throw(_("不支持的作业单据。"))
	if dimension not in ("order", "product"):
		frappe.throw(_("不支持的分拣单维度。"))
	batch = _validate_batch_access(import_batch)
	orders = workflow._batch_orders(batch.name)
	for po in orders:
		po.check_permission("read")
	if purchase_order and purchase_order not in {po.name for po in orders}:
		frappe.throw(_("采购单不属于当前批次。"))
	draft = _draft_rows(draft_rows)
	if stage == "stocking":
		sections = _stocking_sections(batch.name, orders, draft)
	elif stage == "assembly":
		sections = _assembly_sections(batch.name)
	elif dimension == "product":
		sections = _sorting_product_sections(batch.name, orders, purchase_order, draft)
	else:
		sections = _sorting_sections(batch.name, orders, purchase_order, draft)
	for section in sections:
		total = sum(c["width"] for c in section["columns"])
		for c in section["columns"]:
			c["width_percent"] = c["width"] * 100 / total
	return {"title": "商品分拣矩阵单" if stage == "sorting" and dimension == "product" else TITLES[stage], "dimension": dimension, "batch": batch, "sections": sections,
		"printed_at": format_datetime(now_datetime()),
		"source_note": "当前页面草稿（未保存）" if draft is not None else "已保存数据 / 需求清单"}


def _stocking_sections(import_batch, orders, draft):
	data = workflow.get_stocking_workbench(import_batch)
	pools = data["pool_rows"]
	if draft is not None:
		allowed = {r["stock_item"] for r in data["requirements"]}
		if any(r.get("stock_item") not in allowed for r in draft):
			frappe.throw(_("备货明细包含本批次以外的物料。"))
		pools = draft
	links = defaultdict(set)
	for po in orders:
		for item in po.items:
			if item.platform_item:
				for comp in workflow.expand_platform_item(item.platform_item, item.purchase_qty):
					links[comp.stock_item].add(item.jd_sku)
	rows = []
	for req in data["requirements"]:
		matched = [r for r in pools if r.get("stock_item") == req["stock_item"]] or [{}]
		for index, source in enumerate(matched):
			qty = source.get("stocked_qty")
			# Finished bundles have consumed component pools; print original confirmed picking quantity.
			if draft is None and source:
				qty = flt(source.get("required_qty")) - flt(source.get("shortage_qty"))
			rows.append({"stock_item": req["stock_item"], "name": req["stock_item_name"],
				"goods": "\n".join(sorted(links[req["stock_item"]])), "required": req["required_qty"] if index == 0 else None,
				"uom": req["uom"], "warehouse": source.get("warehouse") or "", "batch_no": source.get("batch_no") or "",
				"qty": qty, "remarks": ""})
	for item in data.get("unmapped_requirements", []):
		rows.append({"stock_item": "未映射", "name": item.get("jd_item_name", ""), "goods": item.get("jd_sku", ""), "remarks": "请维护商品映射"})
	columns = [column("stock_item", "ERP物料编码", 115), column("name", "物料名称", 210), column("goods", "关联京东SKU / 主商品", 180), column("required", "需求数量", 65, True), column("uom", "单位", 40), column("warehouse", "备货仓库", 110), column("batch_no", "备货批号", 100), column("qty", "备货数量", 65, True), column("remarks", "备注", 70)]
	return [{"heading": "实物物料备货清单", "note": "相同物料按仓库、批号分行；需求数量仅在首行列出。未维护批号时留空。", "columns": columns, "rows": rows}]


def _assembly_sections(import_batch):
	data = workflow.get_assembly_workbench(import_batch)
	pools = workflow.get_stocking_workbench(import_batch)["pool_rows"]
	rows = []
	for group in data["assembly_rows"]:
		rows.append({"kind": "主商品", "code": group["platform_item"], "name": group["jd_item_name"] or group["stock_item_name"], "uom": group["uom"], "required": group["required_qty"], "done": group["assembled_qty"], "status": "已确认" if group["done"] else "待确认", "_parent": True})
		for comp in group["components"]:
			batches = sorted({r.batch_no for r in pools if r.stock_item == comp["stock_item"] and r.batch_no})
			rows.append({"kind": "子物料", "code": comp["stock_item"], "name": comp["stock_item_name"], "uom": comp["uom"], "ratio": comp["per_set"], "required": comp["per_set"] * group["required_qty"], "batch_no": "、".join(batches)})
	columns = [column("kind", "类型", 60), column("code", "ERP编码", 120), column("name", "商品 / 物料名称", 300), column("uom", "单位", 45), column("ratio", "每套用量", 70, True), column("required", "需求数量", 75, True), column("done", "已组装", 75, True), column("batch_no", "备货批号", 150), column("status", "确认状态", 75), column("remarks", "实作 / 备注", 100)]
	return [{"heading": "主商品及子物料", "note": "子物料紧随所属主商品列示。", "columns": columns, "rows": rows}]


def _sorting_data(import_batch, draft):
	data = workflow.get_sorting_matrix(import_batch)
	allocations = data["allocations"]
	if draft is not None:
		pools = {r.name: r for r in data["pools"]}
		keys = {(r["purchase_order"], r["jd_sku"], r["stock_item"]) for r in data["requirements"]}
		allocations = []
		for row in draft:
			pool = pools.get(row.get("stocking_pool_item"))
			if not pool or (row.get("purchase_order"), row.get("jd_sku"), row.get("stock_item")) not in keys or pool.stock_item != row.get("stock_item"):
				frappe.throw(_("分拣明细不属于当前批次。"))
			if flt(row.get("sorted_qty")) > 0:
				allocations.append({**row, "warehouse": pool.warehouse, "batch_no": pool.batch_no})
	return data, allocations


def _sorting_sections(import_batch, orders, purchase_order, draft):
	data, allocations = _sorting_data(import_batch, draft)
	sections = []
	columns = [column("sku", "京东SKU", 120), column("code", "ERP编码", 130), column("name", "商品名称", 270), column("required", "需求数量", 75, True), column("qty", "分拣数量", 75, True), column("uom", "单位", 45), column("warehouse", "备货仓库", 140), column("batch_no", "发货批号", 110), column("remarks", "备注", 110)]
	for po in orders:
		if purchase_order and po.name != purchase_order:
			continue
		rows = []
		for req in [r for r in data["requirements"] if r["purchase_order"] == po.name]:
			item = next((i for i in po.items if i.jd_sku == req["jd_sku"]), None)
			matching = [r for r in allocations if r.get("purchase_order") == po.name and r.get("jd_sku") == req["jd_sku"] and r.get("stock_item") == req["stock_item"]] or [{}]
			remaining = max(flt(req["required_qty"]) - sum(flt(r.get("sorted_qty")) for r in matching), 0)
			is_bundle = workflow.is_product_bundle(req["platform_item"])
			uom = item.purchase_uom if item and is_bundle else workflow._item_display(req["stock_item"])["uom"]
			component_batches = ""
			if is_bundle:
				components = {c.stock_item for c in workflow.expand_platform_item(req["platform_item"], 1)}
				component_batches = "；".join(sorted({f"{p.stock_item}: {p.batch_no}" for p in data["pools"] if p.stock_item in components and p.batch_no}))
			for index, row in enumerate(matching):
				rows.append({"sku": req["jd_sku"], "code": req["platform_item"], "name": item.jd_item_name if item else "", "uom": uom or "", "required": req["required_qty"] if index == 0 else None, "qty": row.get("sorted_qty"), "warehouse": row.get("warehouse", ""), "batch_no": row.get("batch_no") or component_batches, "remarks": f"待分拣 {remaining:g}" if index == 0 and remaining else ("批号为子物料备货批号" if component_batches else "")})
		sections.append({"heading": f"采购单：{po.name}　目的城市：{po.destination_city or ''}", "note": f"目标仓：{po.jd_warehouse or ''}　已生成箱数：{po.total_cartons or 0}（以实际箱记录为准）", "columns": [dict(c) for c in columns], "rows": rows})
	combined = []
	for section in sections:
		combined.append({"_group": section["heading"] + "　" + section["note"]})
		combined.extend(section["rows"])
	return [{"heading": "采购单 / 目标仓分组明细", "note": "全部采购单连续显示；需求数量只在同一商品的首行列出。", "columns": columns, "rows": combined}]


def _sorting_product_sections(import_batch, orders, purchase_order, draft):
	data, allocations = _sorting_data(import_batch, draft)
	visible = [po for po in orders if not purchase_order or po.name == purchase_order]
	columns = [column("product", "ERP物料 / 京东SKU / 批号", 220), column("pool_qty", "池数量", 65, True)]
	columns += [column(f"po_{i}", f"{po.name}\n{po.distribution_center or po.destination_city or ''}", 110) for i, po in enumerate(visible)]
	requirements = {(r["purchase_order"], r["jd_sku"], r["stock_item"]): r for r in data["requirements"]}
	allocated = defaultdict(float)
	for row in allocations:
		allocated[(row.get("stocking_pool_item"), row.get("purchase_order"), row.get("jd_sku"), row.get("stock_item"))] += flt(row.get("sorted_qty"))
	rows = []
	for pool in data["pools"]:
		skus = sorted({r["jd_sku"] for r in data["requirements"] if r["stock_item"] == pool.stock_item})
		for sku in skus:
			row = {"product": f"{pool.stock_item}\n{sku} / {pool.batch_no or '—'}", "pool_qty": pool.stocked_qty}
			for i, po in enumerate(visible):
				req = requirements.get((po.name, sku, pool.stock_item))
				qty = allocated[(pool.name, po.name, sku, pool.stock_item)]
				row[f"po_{i}"] = f"{qty:g}" if req else "—"
			rows.append(row)
	return [{"heading": "按商品分拣矩阵", "note": "与页面一致：行是物料 / SKU / 批号，列是采购单；单元格仅显示分拣数量。池数量按原备货行展示，请勿跨 SKU 重复汇总。", "columns": columns, "rows": rows}]


def render_report(report):
	return frappe.render_template("yimed_ecommerce/templates/jd_workbench_report.html", {**report, "format_cell": _format_cell})


def report_xlsx(report):
	from openpyxl import Workbook
	from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
	from openpyxl.utils import get_column_letter
	book = Workbook()
	book.remove(book.active)
	for index, section in enumerate(report["sections"] or [{"heading": "暂无数据", "note": "", "columns": [column("empty", "暂无数据")], "rows": []}]):
		sheet = book.create_sheet(f"清单{index + 1}")
		cols = section["columns"]
		sheet.append([report["title"]]); sheet.merge_cells(start_row=1, start_column=1, end_row=1, end_column=len(cols))
		sheet.append([f"公司：{report['batch'].company}　批次：{report['batch'].name}"])
		sheet.append([section["heading"]]); sheet.append([section["note"]])
		for n in (2, 3, 4): sheet.merge_cells(start_row=n, start_column=1, end_row=n, end_column=len(cols))
		sheet.append([c["label"] for c in cols])
		for row in section["rows"]:
			# Explicit text cells prevent formula execution and preserve long SKU identifiers.
			if row.get("_group"):
				sheet.append([row["_group"]])
				sheet.merge_cells(start_row=sheet.max_row, start_column=1, end_row=sheet.max_row, end_column=len(cols))
			else:
				sheet.append([row.get(c["fieldname"]) for c in cols])
			for cell in sheet[sheet.max_row]:
				if isinstance(cell.value, str): cell.data_type = "s"
			if row.get("_parent") or row.get("_group"):
				for cell in sheet[sheet.max_row]: cell.fill = PatternFill("solid", fgColor="EEF2F5")
		for cells in sheet:
			for cell in cells:
				cell.font = Font(name="Microsoft YaHei", size=10, bold=cell.row in (1, 5))
				cell.alignment = Alignment(vertical="center", wrap_text=True)
				if cell.row >= 5: cell.border = Border(bottom=Side(style="thin", color="B0B0B0"))
		for j, c in enumerate(cols, 1): sheet.column_dimensions[get_column_letter(j)].width = max(9, c["width"] / 7)
		sheet.freeze_panes = "A6"
		if not any(row.get("_group") for row in section["rows"]):
			sheet.auto_filter.ref = f"A5:{get_column_letter(len(cols))}{sheet.max_row}"
		sheet.sheet_properties.pageSetUpPr.fitToPage = True
		sheet.page_setup.orientation = "landscape"; sheet.page_setup.paperSize = sheet.PAPERSIZE_A4
		sheet.page_setup.fitToWidth = 1; sheet.page_setup.fitToHeight = 0
		sheet.print_title_rows = "1:5"
		sheet.oddFooter.center.text = "第 &P 页 / 共 &N 页"
		sheet.oddFooter.left.text = report["source_note"]
	out = BytesIO(); book.save(out); return out.getvalue()


def report_pdf(report):
	from pypdf import PdfReader, PdfWriter, PageObject
	from pypdf.generic import DictionaryObject, NameObject, DecodedStreamObject
	if report.get("dimension") == "product":
		sections = []
		for section in report["sections"]:
			fixed, orders = section["columns"][:2], section["columns"][2:]
			for offset in range(0, max(len(orders), 1), 5):
				cols = [dict(c) for c in fixed + orders[offset:offset + 5]]
				total = sum(c["width"] for c in cols)
				for c in cols: c["width_percent"] = c["width"] * 100 / total
				sections.append({**section, "columns": cols, "heading": f"按商品分拣矩阵（采购单列 {offset + 1}–{min(offset + 5, len(orders))}）"})
		report = {**report, "sections": sections}
	reader = PdfReader(BytesIO(_render_chrome_pdf(render_report(report))))
	writer = PdfWriter()
	for index, page in enumerate(reader.pages, 1):
		width, height = float(page.mediabox.width), float(page.mediabox.height)
		overlay = PageObject.create_blank_page(width=width, height=height)
		font = DictionaryObject({NameObject("/Type"): NameObject("/Font"), NameObject("/Subtype"): NameObject("/Type1"), NameObject("/BaseFont"): NameObject("/Helvetica")})
		overlay[NameObject("/Resources")] = DictionaryObject({NameObject("/Font"): DictionaryObject({NameObject("/PageNumberFont"): font})})
		stream = DecodedStreamObject()
		stream.set_data(f"q BT /PageNumberFont 9 Tf 0.35 g 1 0 0 1 {width / 2 - 12} 16 Tm ({index} / {len(reader.pages)}) Tj ET Q".encode("ascii"))
		overlay[NameObject("/Contents")] = stream
		page.merge_page(overlay)
		writer.add_page(page)
	out = BytesIO(); writer.write(out); return out.getvalue()


@frappe.whitelist()
def download_workbench_report(import_batch: str, stage: str, output: str = "pdf", purchase_order: str | None = None, draft_rows=None, dimension: str = "order"):
	if output not in ("pdf", "xlsx"):
		frappe.throw(_("不支持的导出格式。"))
	report = build_report(import_batch, stage, purchase_order, draft_rows, dimension)
	content = report_pdf(report) if output == "pdf" else report_xlsx(report)
	frappe.local.response.filename = f"{report['title']}-{import_batch}.{output}"
	frappe.local.response.filecontent = content
	frappe.local.response.type = "pdf" if output == "pdf" else "binary"
	frappe.local.response.display_content_as = "inline" if output == "pdf" else "attachment"
