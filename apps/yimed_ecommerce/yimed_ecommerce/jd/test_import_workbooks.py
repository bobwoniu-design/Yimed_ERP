from io import BytesIO

import frappe
from frappe.tests import IntegrationTestCase
from openpyxl import load_workbook

from yimed_ecommerce.jd.import_workbooks import (
	IMPORT_COLUMNS,
	build_failed_details,
	build_import_template,
	download_failed_details,
	download_import_template,
)


class TestJDImportWorkbooks(IntegrationTestCase):
	def test_import_template_contains_upload_sheet_and_instructions(self):
		workbook = load_workbook(BytesIO(build_import_template()), data_only=False)
		self.assertEqual(workbook.sheetnames, ["采购订单导入", "填写说明"])
		data_sheet = workbook["采购订单导入"]
		self.assertEqual(
			[data_sheet.cell(1, column).value for column in range(1, len(IMPORT_COLUMNS) + 1)],
			[column[0] for column in IMPORT_COLUMNS],
		)
		instructions = workbook["填写说明"]
		self.assertIn("可以直接上传京东后台导出的采购订单明细", instructions["A1"].value)
		required = {
			instructions.cell(row, 1).value
			for row in range(4, instructions.max_row + 1)
			if instructions.cell(row, 2).value == "必填"
		}
		self.assertEqual(required, {"采购单号", "商品编号", "采购数量"})

	def test_failed_details_are_fixed_text_not_excel_formulas(self):
		batch = frappe._dict(
			{
				"error_log": "=2+3: =HYPERLINK(\"https://invalid.example\",\"click\")",
			}
		)
		workbook = load_workbook(BytesIO(build_failed_details(batch)), data_only=False)
		sheet = workbook["失败明细"]
		self.assertEqual(sheet["A2"].value, "=2+3")
		self.assertEqual(sheet["B2"].value, '=HYPERLINK("https://invalid.example","click")')
		self.assertEqual(sheet["A2"].data_type, "s")
		self.assertEqual(sheet["B2"].data_type, "s")

	def test_failed_details_refuses_empty_error_log(self):
		with self.assertRaises(frappe.ValidationError):
			build_failed_details(frappe._dict({"error_log": ""}))

	def test_download_template_requires_batch_read_permission(self):
		original_user = frappe.session.user
		try:
			frappe.set_user("Guest")
			with self.assertRaises(frappe.PermissionError):
				download_import_template()
		finally:
			frappe.set_user(original_user)

	def test_download_failed_details_requires_batch_read_permission(self):
		batch = frappe.get_doc(
			{
				"doctype": "JD Purchase Import Batch",
				"company": frappe.db.get_value("Company", {}, "name"),
				"import_file": "/private/files/permission-test.xlsx",
				"failed_count": 1,
				"error_log": "PO-1: Test failure",
			}
		).insert()
		original_user = frappe.session.user
		try:
			frappe.set_user("Guest")
			with self.assertRaises(frappe.PermissionError):
				download_failed_details(batch.name)
		finally:
			frappe.set_user(original_user)
