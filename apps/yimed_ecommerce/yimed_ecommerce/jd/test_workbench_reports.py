import unittest
from io import BytesIO
from types import SimpleNamespace
from unittest.mock import patch

from openpyxl import load_workbook
from yimed_ecommerce.jd import workbench_reports as reports


class TestWorkbenchReports(unittest.TestCase):
	def setUp(self):
		self.frappe = self.enterContext(patch.object(reports, "frappe"))
		self.enterContext(patch.object(reports, "_", lambda text: text))
		self.frappe.throw.side_effect = ValueError("Invalid report")
		self.data = {"requirements": [{"stock_item": "ITEM", "stock_item_name": "测试物料", "required_qty": 10, "uom": "个"}], "pool_rows": [], "unmapped_requirements": []}
		self.enterContext(patch.object(reports.workflow, "get_stocking_workbench", return_value=self.data))

	def test_unconfirmed_stocking_without_batch_is_printable(self):
		row = reports._stocking_sections("BATCH", [], None)[0]["rows"][0]
		self.assertEqual(row["batch_no"], "")
		self.assertEqual(row["required"], 10)

	def test_unsaved_batches_are_preserved_in_separate_rows(self):
		draft = [{"stock_item": "ITEM", "batch_no": "B1", "stocked_qty": 4}, {"stock_item": "ITEM", "batch_no": "B2", "stocked_qty": 6}]
		rows = reports._stocking_sections("BATCH", [], draft)[0]["rows"]
		self.assertEqual([r["batch_no"] for r in rows], ["B1", "B2"])
		self.assertEqual([r["qty"] for r in rows], [4, 6])
		self.assertEqual([r["required"] for r in rows], [10, None])

	def test_assembled_components_print_original_picking_quantity(self):
		self.data["pool_rows"] = [{"stock_item": "ITEM", "required_qty": 10, "shortage_qty": 2, "stocked_qty": 0, "batch_no": "B1"}]
		row = reports._stocking_sections("BATCH", [], None)[0]["rows"][0]
		self.assertEqual(row["qty"], 8)

	def test_foreign_draft_item_is_rejected(self):
		with self.assertRaises(ValueError):
			reports._stocking_sections("BATCH", [], [{"stock_item": "FOREIGN"}])

	def test_foreign_purchase_order_is_rejected_after_access_check(self):
		with patch.object(reports, "_validate_batch_access", return_value=SimpleNamespace(name="BATCH")) as access, patch.object(reports.workflow, "_batch_orders", return_value=[]):
			with self.assertRaises(ValueError): reports.build_report("BATCH", "sorting", "FOREIGN")
			access.assert_called_once_with("BATCH")

	def test_excel_keeps_identifiers_and_formula_like_text_as_text(self):
		report = {"title": "备货单", "batch": SimpleNamespace(company="公司", name="BATCH"), "source_note": "草稿", "sections": [{"heading": "明细", "note": "", "columns": [reports.column("code", "商品编码"), reports.column("name", "名称")], "rows": [{"code": "001234567890123456", "name": "=1+1"}]}]}
		book = load_workbook(BytesIO(reports.report_xlsx(report)))
		self.assertEqual(book.active["A6"].value, "001234567890123456")
		self.assertEqual(book.active["B6"].data_type, "s")
		self.assertEqual(book.active.page_setup.orientation, "landscape")
		self.assertEqual(str(book.active.page_setup.paperSize), str(book.active.PAPERSIZE_A4))

	def test_product_matrix_matches_pool_and_order_axes_without_warehouse(self):
		orders = [SimpleNamespace(name="PO1", distribution_center="", destination_city="武汉"), SimpleNamespace(name="PO2", distribution_center="", destination_city="成都")]
		data = {"pools": [SimpleNamespace(name="P1", stock_item="ITEM", batch_no="B1", stocked_qty=10)], "requirements": [{"purchase_order": "PO1", "jd_sku": "SKU", "stock_item": "ITEM", "required_qty": 6}, {"purchase_order": "PO2", "jd_sku": "SKU", "stock_item": "ITEM", "required_qty": 4}]}
		allocations = [{"stocking_pool_item": "P1", "purchase_order": "PO1", "jd_sku": "SKU", "stock_item": "ITEM", "sorted_qty": 6}]
		with patch.object(reports, "_sorting_data", return_value=(data, allocations)):
			section = reports._sorting_product_sections("BATCH", orders, None, None)[0]
			self.assertEqual(len(section["columns"]), 4)
			self.assertEqual(section["rows"][0]["po_0"], "6")
			self.assertEqual(section["rows"][0]["po_1"], "0")
			self.assertNotIn("warehouse", [c["fieldname"] for c in section["columns"]])
			filtered = reports._sorting_product_sections("BATCH", orders, "PO2", None)[0]
			self.assertEqual(len(filtered["columns"]), 3)
			self.assertTrue(filtered["columns"][-1]["label"].startswith("PO2"))

	def test_order_groups_share_a_single_worksheet(self):
		orders = [SimpleNamespace(name=name, destination_city="武汉", jd_warehouse="目标仓", total_cartons=0, items=[]) for name in ("PO1", "PO2")]
		with patch.object(reports, "_sorting_data", return_value=({"requirements": []}, [])):
			sections = reports._sorting_sections("BATCH", orders, None, None)
		self.assertEqual(len(sections), 1)
		self.assertEqual(len(sections[0]["rows"]), 2)
		report = {"title": "采购订单分拣单", "batch": SimpleNamespace(company="公司", name="BATCH"), "source_note": "已保存", "sections": sections}
		book = load_workbook(BytesIO(reports.report_xlsx(report)))
		self.assertEqual(len(book.sheetnames), 1)
		self.assertIn("PO1", book.active["A6"].value)
		self.assertIn("PO2", book.active["A7"].value)
